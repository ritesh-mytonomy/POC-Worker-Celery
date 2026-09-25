"""Build every test fixture in design.md §10.1 into fixtures/out/ (or a given directory).

Usage: python fixtures/make_fixtures.py [OUT_DIR]
"""
import io
import random
import struct
import sys
import warnings
import zipfile
from collections.abc import Callable
from pathlib import Path

from docx import Document

DATE = (2026, 9, 24, 12, 0, 0)           # fixed timestamps so the zip wrappers are stable
BOMB_BYTES = 200 * 1024 * 1024           # one 200 MB entry of zeros
LYING_REAL_BYTES = 256 * 1024            # real size of lying.zip's entry
LYING_DECLARED_BYTES = 100 * 1024        # size its central directory claims
CENTRAL_SIG = b"PK\x01\x02"


def docx_bytes(text: str) -> bytes:
    """A valid .docx (python-docx) with one paragraph."""
    buf = io.BytesIO()
    document = Document()
    document.add_paragraph(text)
    document.save(buf)
    return buf.getvalue()


def _info(name: str, compress: int = zipfile.ZIP_DEFLATED) -> zipfile.ZipInfo:
    """A ZipInfo with a fixed timestamp."""
    info = zipfile.ZipInfo(name, date_time=DATE)
    info.compress_type = compress
    return info


def _zip(members: list[tuple[str, bytes]]) -> bytes:
    """A deflated zip of (name, data) members, in order."""
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as zf:
        for name, data in members:
            zf.writestr(_info(name), data)
    return buf.getvalue()


def _central_offset(data: bytes) -> int:
    """Offset of the (single) central directory record."""
    offset = data.rfind(CENTRAL_SIG)
    if offset < 0:
        raise ValueError("no central directory record")
    return offset


def valid_docx() -> bytes:
    """valid.docx — S2."""
    return docx_bytes("A valid ClinSync test document.")


def renamed_exe_docx() -> bytes:
    """renamed_exe.docx — 4 KB beginning MZ (S3)."""
    return b"MZ" + bytes(random.Random(1).randbytes(4096 - 2))


def renamed_zip_docx() -> bytes:
    """renamed_zip.docx — a plain zip of two .txt files (S3)."""
    return _zip([("one.txt", b"first text file\n"), ("two.txt", b"second text file\n")])


def notes_pdf() -> bytes:
    """A minimal PDF, used inside mixed.zip."""
    return (b"%PDF-1.4\n1 0 obj << /Type /Catalog /Pages 2 0 R >> endobj\n"
            b"2 0 obj << /Type /Pages /Kids [] /Count 0 >> endobj\ntrailer << /Root 1 0 R >>\n%%EOF\n")


def mixed_zip() -> bytes:
    """mixed.zip — 3 valid .docx + notes.pdf + readme.txt (S1)."""
    return _zip([(f"doc{i}.docx", docx_bytes(f"Mixed document {i}")) for i in range(1, 4)]
                + [("notes.pdf", notes_pdf()), ("readme.txt", b"Read me.\n")])


def big30_zip() -> bytes:
    """big30.zip — 30 valid .docx (S4, S5)."""
    return _zip([(f"doc{i:02d}.docx", docx_bytes(f"Document {i} of 30")) for i in range(1, 31)])


def bomb_zip() -> bytes:
    """bomb.zip — one 200 MB entry of zeros, deflated (~1000:1)."""
    buf = io.BytesIO()
    zeros = bytes(1024 * 1024)
    with zipfile.ZipFile(buf, "w") as zf:
        with zf.open(_info("zeros.bin"), "w", force_zip64=True) as dst:
            for _ in range(BOMB_BYTES // len(zeros)):
                dst.write(zeros)
    return buf.getvalue()


def slip_zip() -> bytes:
    """slip.zip — one entry named ../../evil.docx."""
    return _zip([("../../evil.docx", docx_bytes("path traversal"))])


def encrypted_zip() -> bytes:
    """encrypted.zip — one entry with flag_bits |= 0x1 in both the local header and the central directory."""
    data = bytearray(_zip([("secret.docx", docx_bytes("encrypted"))]))
    for flag_at in (6, _central_offset(bytes(data)) + 8):          # general-purpose bit flag, both headers
        (flags,) = struct.unpack_from("<H", data, flag_at)
        struct.pack_into("<H", data, flag_at, flags | 0x1)
    return bytes(data)


def lying_zip() -> bytes:
    """lying.zip — one entry whose central-directory uncompressed size is patched below its real size."""
    real = random.Random(2).randbytes(LYING_REAL_BYTES // 2) * 2      # compressible enough to stay under ratio
    data = bytearray(_zip([("report.docx", real)]))
    struct.pack_into("<I", data, _central_offset(bytes(data)) + 24, LYING_DECLARED_BYTES)
    return bytes(data)


def dupnames_zip() -> bytes:
    """dupnames.zip — two entries with the same name, a.docx (rev 1.3, R6.9)."""
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as zf:
        with warnings.catch_warnings():
            warnings.simplefilter("ignore", UserWarning)             # zipfile warns on duplicate names
            zf.writestr(_info("a.docx"), docx_bytes("first a.docx"))
            zf.writestr(_info("a.docx"), docx_bytes("second a.docx"))
    return buf.getvalue()


def lying3_zip() -> bytes:
    """lying3.zip — three valid .docx, stored; the third's central-directory sizes patched to half (task 9.3)."""
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as zf:
        for i in range(3):
            zf.writestr(_info(f"d{i}.docx", zipfile.ZIP_STORED), docx_bytes(f"Lying archive document {i}"))
    data = bytearray(buf.getvalue())
    third = _central_offset(bytes(data))
    (real,) = struct.unpack_from("<I", data, third + 24)
    struct.pack_into("<I", data, third + 20, real // 2)      # compressed size (stored)
    struct.pack_into("<I", data, third + 24, real // 2)      # uncompressed size
    return bytes(data)


def inner_bomb_zip() -> bytes:
    """inner_bomb.zip — a.docx, bomb.docx (bomb.zip's bytes, stored so the outer ratio stays ~1), c.docx (task 9.3)."""
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as zf:
        zf.writestr(_info("a.docx"), docx_bytes("Inner bomb archive, first"))
        zf.writestr(_info("bomb.docx", zipfile.ZIP_STORED), bomb_zip())
        zf.writestr(_info("c.docx"), docx_bytes("Inner bomb archive, last"))
    return buf.getvalue()


def docx_of_size(target: int, seed: int) -> bytes:
    """A valid .docx of roughly `target` bytes: random words don't compress, so they set the size."""
    rng = random.Random(seed)
    words = ["".join(rng.choice("abcdefghijklmnopqrstuvwxyz") for _ in range(rng.randint(3, 10)))
             for _ in range(2000)]
    document = Document()
    size, paragraphs = 0, 0
    while size < target:
        document.add_paragraph(" ".join(rng.choice(words) for _ in range(400)))
        paragraphs += 1
        if paragraphs % 5 == 0:
            buf = io.BytesIO()
            document.save(buf)
            size = len(buf.getvalue())
    buf = io.BytesIO()
    document.save(buf)
    return buf.getvalue()


def big100_zip() -> bytes:
    """big100.zip — 100 valid .docx of ~50 KB each (NFR-3)."""
    doc = docx_of_size(50 * 1024, seed=3)
    return _zip([(f"doc{i:03d}.docx", doc) for i in range(1, 101)])


FIXTURES: dict[str, Callable[[], bytes]] = {
    "valid.docx": valid_docx,
    "renamed_exe.docx": renamed_exe_docx,
    "renamed_zip.docx": renamed_zip_docx,
    "mixed.zip": mixed_zip,
    "big30.zip": big30_zip,
    "bomb.zip": bomb_zip,
    "slip.zip": slip_zip,
    "encrypted.zip": encrypted_zip,
    "lying.zip": lying_zip,
    "dupnames.zip": dupnames_zip,
    "lying3.zip": lying3_zip,
    "inner_bomb.zip": inner_bomb_zip,
    "big100.zip": big100_zip,
}


def build(out_dir: Path) -> list[Path]:
    """Write every fixture into out_dir; return their paths."""
    out_dir.mkdir(parents=True, exist_ok=True)
    paths = []
    for name, make in FIXTURES.items():
        path = out_dir / name
        path.write_bytes(make())
        paths.append(path)
    return paths


def main() -> None:
    """Build into fixtures/out/, or the directory given as the first argument."""
    out_dir = Path(sys.argv[1]) if len(sys.argv) > 1 else Path(__file__).parent / "out"
    for path in build(out_dir):
        print(f"{path.stat().st_size:>11,}  {path}")


if __name__ == "__main__":
    main()
